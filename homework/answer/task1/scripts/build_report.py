"""Build the report using the old report's restrained typography.

Use bundled Codex Python. Teacher geometry and existing package parts stay intact.
Markdown: # / ## / ###, **bold**, `inline code`, simple pipe tables, image groups,
<!-- page -->, <!-- images: name.png | name.png -->, <!-- code: name.png -->.
Image names resolve in docs/screenshots; code names in docs/code-screenshots.
"""
from __future__ import annotations

from copy import deepcopy
from hashlib import sha256
from io import BytesIO
import json
from pathlib import Path
import re
from zipfile import ZipFile, ZIP_DEFLATED

from docx import Document
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor
from lxml import etree
from PIL import Image

TASK = Path(__file__).resolve().parents[1]
REPO = TASK.parents[2]
REFERENCE = REPO / "homework/teacher/task1/作业报告.docx"
FINAL = TASK / "docs/作业报告.docx"
MARKDOWN = TASK / "docs/报告.md"
QA = TASK / "verification/report"
REFERENCE_HASH = "cd682fbbda36cce208e9661699b5353fe4429ba343daa9787da6a6793ad24e16"
STYLES = {"Title", "ReportSubtitle", "ReportAuthor", "ReportBody",
          "ReportHeading", "ReportSubheading", "ReportCaption", "ReportTable", "ReportList"}


def digest(data: bytes) -> str:
    return sha256(data).hexdigest()


def append_xml(original: bytes, children: list[bytes], closing: bytes) -> bytes:
    if not children:
        return original
    if original.count(closing) != 1:
        raise ValueError(f"Ambiguous XML closing marker {closing!r}")
    return original.replace(closing, b"".join(children) + closing)


def fonts(rpr, latin="Times New Roman", chinese="宋体") -> None:
    rf = rpr.get_or_add_rFonts()
    for name, value in {"ascii": latin, "hAnsi": latin, "cs": latin, "eastAsia": chinese}.items():
        rf.set(qn("w:" + name), value)
    for key in ("asciiTheme", "hAnsiTheme", "csTheme", "eastAsiaTheme"):
        rf.attrib.pop(qn("w:" + key), None)


def inline(p, content: str, *, bold=False) -> None:
    for token in re.split(r"(\*\*.*?\*\*|`[^`]+`)", content):
        if not token:
            continue
        strong = token.startswith("**") and token.endswith("**")
        code = token.startswith("`") and token.endswith("`")
        run = p.add_run(token[2:-2] if strong else token[1:-1] if code else token)
        run.font.bold = bold or strong
        run.font.color.rgb = RGBColor(0, 0, 0)
        if code:
            fonts(run._r.get_or_add_rPr(), "Consolas", "宋体")
            run.font.size = Pt(10.5)


def add_styles(doc) -> None:
    specs = {"Title": (24, "黑体", True), "ReportSubtitle": (18, "黑体", False),
             "ReportAuthor": (14, "宋体", False), "ReportBody": (12, "宋体", False),
             "ReportHeading": (14, "黑体", True), "ReportSubheading": (12, "黑体", True),
             "ReportCaption": (11, "宋体", False), "ReportTable": (10.5, "宋体", False),
             "ReportList": (12, "宋体", False)}
    for name, (size, chinese, bold) in specs.items():
        if name in doc.styles:
            raise ValueError(f"Task-specific style unexpectedly exists in template: {name}")
        s = doc.styles.add_style(name, WD_STYLE_TYPE.PARAGRAPH)
        s.base_style = doc.styles["Normal"]
        s.font.name = "Times New Roman"
        s.font.size = Pt(size)
        s.font.bold = bold
        s.font.color.rgb = RGBColor(0, 0, 0)
        fonts(s.element.get_or_add_rPr(), "Times New Roman", chinese)
        pf = s.paragraph_format
        pf.line_spacing_rule = WD_LINE_SPACING.EXACTLY
        pf.line_spacing = Pt(20 if size <= 14 else size * 1.4)
        pf.space_after = Pt(6)
        pf.widow_control = True
        pf.first_line_indent = Pt(24 if name == "ReportBody" else 0)
        pf.keep_with_next = name in {"Title", "ReportSubtitle", "ReportAuthor", "ReportHeading", "ReportSubheading"}
        if name in {"ReportHeading", "ReportSubheading"}:
            pf.space_before = Pt(10 if name == "ReportHeading" else 6)
        if name in {"ReportCaption", "ReportTable"}:
            pf.line_spacing = Pt(16)
        if name == "ReportTable":
            pf.space_after = Pt(2)


def table_cells(line: str) -> list[str]:
    return [part.strip().replace(r"\|", "|") for part in re.split(r"(?<!\\)\|", line.strip().strip("|"))]


def build() -> None:
    if digest(REFERENCE.read_bytes()) != REFERENCE_HASH:
        raise ValueError("The retained report template hash changed.")
    text = MARKDOWN.read_text(encoding="utf-8")
    unresolved = re.findall(r"<!-- (?:robot|pixel-scene|texture|validation)-description -->", text)
    if unresolved:
        raise ValueError(f"Report evidence slots remain unresolved: {unresolved}")
    if any(marker in text for marker in ('<!-- SIMD-results -->','<!-- AA-results -->','<!-- GUI-results -->')):
        raise ValueError("Measured extension results are still unresolved.")
    doc = Document(REFERENCE)
    section_before = deepcopy(doc.sections[0]._sectPr)
    body = doc._element.body
    for element in list(body):
        if element.tag != qn("w:sectPr"):
            body.remove(element)
    add_styles(doc)
    pending_break = False
    figures, table_audit = [], []
    code_manifest_path = TASK / "verification/code-screenshots.json"
    code_manifest = (json.loads(code_manifest_path.read_text(encoding="utf-8"))
                     if code_manifest_path.is_file() else {})
    page_markers = 0
    body_width = (doc.sections[0].page_width - doc.sections[0].left_margin - doc.sections[0].right_margin) / 914400.0

    def paragraph(content="", role="body"):
        nonlocal pending_break
        style = {"title": "Title", "subtitle": "ReportSubtitle", "author": "ReportAuthor",
                 "body": "ReportBody", "heading": "ReportHeading", "subheading": "ReportSubheading",
                 "caption": "ReportCaption", "list": "ReportList"}[role]
        p = doc.add_paragraph(style=style)
        if pending_break:
            p.paragraph_format.page_break_before = True
            pending_break = False
        if role in {"title", "subtitle", "author", "caption"}:
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        elif role == "body":
            # Long URLs and reproduction commands produce stretched short lines
            # under Word justification; use normal left alignment for these.
            p.alignment = (WD_ALIGN_PARAGRAPH.LEFT
                           if re.search(r"https?://|scripts/", content)
                           else WD_ALIGN_PARAGRAPH.JUSTIFY)
        inline(p, content, bold=role in {"title", "heading", "subheading"})
        return p

    def add_table(rows):
        nonlocal pending_break
        headings = table_cells(rows[0])
        alignments = table_cells(rows[1])
        values = [headings] + [table_cells(row) for row in rows[2:]]
        if any(len(row) != len(headings) for row in values):
            raise ValueError("Markdown table has inconsistent column counts.")
        table = doc.add_table(rows=len(values), cols=len(headings))
        table.autofit = False
        weights = [max(6, min(24, max(len(row[c]) for row in values))) for c in range(len(headings))]
        total = sum(weights)
        for c, weight in enumerate(weights):
            table.columns[c].width = Inches(body_width * weight / total)
        borders = OxmlElement("w:tblBorders")
        for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
            node = OxmlElement("w:" + edge)
            for attr, value in {"val": "single", "sz": "4", "color": "D9D9D9"}.items():
                node.set(qn("w:" + attr), value)
            borders.append(node)
        table._tbl.tblPr.append(borders)
        margins = OxmlElement("w:tblCellMar")
        for edge in ("top", "left", "bottom", "right"):
            node = OxmlElement("w:" + edge)
            node.set(qn("w:w"), "70" if edge in {"top", "bottom"} else "90")
            node.set(qn("w:type"), "dxa")
            margins.append(node)
        table._tbl.tblPr.append(margins)
        for r, row in enumerate(table.rows):
            if r == 0:
                row._tr.get_or_add_trPr().append(OxmlElement("w:tblHeader"))
            for c, cell in enumerate(row.cells):
                cell.width = Inches(body_width * weights[c] / total)
                cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
                p = cell.paragraphs[0]
                p.style = doc.styles["ReportTable"]
                marker = alignments[c]
                p.alignment = (WD_ALIGN_PARAGRAPH.CENTER if marker.startswith(":") and marker.endswith(":")
                               else WD_ALIGN_PARAGRAPH.RIGHT if marker.endswith(":") else WD_ALIGN_PARAGRAPH.LEFT)
                if pending_break and r == 0 and c == 0:
                    p.paragraph_format.page_break_before = True
                    pending_break = False
                inline(p, values[r][c], bold=(r == 0))
                if r == 0:
                    shading = OxmlElement("w:shd")
                    shading.set(qn("w:fill"), "F2F2F2")
                    cell._tc.get_or_add_tcPr().append(shading)
        paragraph().paragraph_format.space_after = Pt(3)
        table_audit.append({"rows": len(values), "columns": len(headings), "headings": headings})

    def add_images(spec, is_code=False):
        width_limit = None
        if spec.startswith("width="):
            setting, spec = spec.split(";", 1)
            width_limit = float(setting.split("=", 1)[1])
        names = [item.strip() for item in spec.split("|")]
        p = paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.first_line_indent = Pt(0)
        p.paragraph_format.keep_with_next = True
        p.paragraph_format.space_after = Pt(3)
        p.paragraph_format.line_spacing_rule = WD_LINE_SPACING.SINGLE
        width = (body_width - .1 * (len(names) - 1)) / len(names)
        if width_limit:
            width = min(width, width_limit)
        folder = TASK / "docs" / ("code-screenshots" if is_code else "screenshots")
        for i, name in enumerate(names):
            path = folder / name
            if not path.is_file():
                raise FileNotFoundError(path)
            if is_code:
                relative = str(path.relative_to(TASK)).replace("\\", "/")
                matches = [(entry, image) for entry in code_manifest.get("excerpts", [])
                           for image in entry["images"] if image["file"] == relative]
                if len(matches) != 1:
                    raise ValueError(f"Code image needs a unique source manifest entry: {name}")
                entry, evidence = matches[0]
                if digest((TASK / entry["source"]).read_bytes()) != entry["source_sha256"]:
                    raise ValueError(f"Source changed since code capture; rerun capture_code.py: {entry['source']}")
                if digest(path.read_bytes()) != evidence["sha256"]:
                    raise ValueError(f"Code image hash changed: {name}")
            with Image.open(path) as img:
                if img.format != "PNG":
                    raise ValueError(f"Expected a lossless PNG: {path}")
                pixels = img.size
            shape = p.add_run().add_picture(str(path), width=Inches(width))
            shape._inline.docPr.set("descr", name)
            figures.append({"file": str(path.relative_to(TASK)), "pixels": pixels,
                            "width_inches": width, "kind": "code" if is_code else "render",
                            "sha256": digest(path.read_bytes())})
            if i + 1 < len(names):
                p.add_run(" ")

    lines, buffer = text.splitlines(), []
    title_passed, title_blocks = False, 0

    def flush():
        nonlocal title_blocks
        if not buffer:
            return
        content = "".join(buffer)
        buffer.clear()
        if title_passed and title_blocks < 2:
            paragraph(content, ["subtitle", "author"][title_blocks])
            title_blocks += 1
        else:
            paragraph(content, "caption" if re.match(r"(?:图|表|代码)\s*\d+(?:\.\d+)?", content) else "body")

    i = 0
    while i < len(lines):
        line = lines[i]
        i += 1
        if not line.strip():
            flush()
        elif line.startswith("### "):
            flush(); paragraph(line[4:], "subheading")
        elif line.startswith("## "):
            flush(); paragraph(line[3:], "heading")
        elif line.startswith("# "):
            flush(); paragraph(line[2:], "title"); title_passed = True
        elif line == "<!-- page -->":
            flush(); pending_break = True; page_markers += 1
        elif re.fullmatch(r"<!-- (?:images|code): (.+) -->", line):
            flush()
            match = re.fullmatch(r"<!-- (images|code): (.+) -->", line)
            add_images(match.group(2), match.group(1) == "code")
        elif line.startswith("|") and i < len(lines) and re.fullmatch(r"[\s|:\-]+", lines[i]):
            flush()
            rows = [line, lines[i]]
            i += 1
            while i < len(lines) and lines[i].startswith("|"):
                rows.append(lines[i]); i += 1
            add_table(rows)
        elif re.match(r"^(?:[-*] |\d+\. )", line):
            flush(); paragraph(line, "list")
        elif line.startswith("<!--"):
            flush()
        else:
            buffer.append(line)
    flush()

    def xml_values(element):
        return (element.tag, tuple(sorted(element.attrib.items())), tuple(xml_values(child) for child in element))
    if xml_values(section_before) != xml_values(doc.sections[0]._sectPr):
        raise AssertionError("Report section geometry changed.")
    stream = BytesIO()
    doc.save(stream)
    stream.seek(0)
    preserve_checks = []
    with ZipFile(REFERENCE) as original, ZipFile(stream) as generated:
        parts = {name: original.read(name) for name in original.namelist()}
        parts["word/document.xml"] = generated.read("word/document.xml")
        styles = etree.fromstring(generated.read("word/styles.xml"))
        new_nodes = [etree.tostring(node) for node in styles if node.get(qn("w:styleId")) in STYLES]
        parts["word/styles.xml"] = append_xml(parts["word/styles.xml"], new_nodes, b"</w:styles>")
        rel_name = "word/_rels/document.xml.rels"
        old_ids = {node.get("Id") for node in etree.fromstring(original.read(rel_name))}
        rels = [etree.tostring(node) for node in etree.fromstring(generated.read(rel_name)) if node.get("Id") not in old_ids]
        parts[rel_name] = append_xml(parts[rel_name], rels, b"</Relationships>")
        ct = "[Content_Types].xml"
        old_keys = {(n.tag, n.get("Extension"), n.get("PartName")) for n in etree.fromstring(original.read(ct))}
        types = [etree.tostring(n) for n in etree.fromstring(generated.read(ct)) if (n.tag, n.get("Extension"), n.get("PartName")) not in old_keys]
        parts[ct] = append_xml(parts[ct], types, b"</Types>")
        for name in generated.namelist():
            if name.startswith("word/media/"):
                parts[name] = generated.read(name)
        editable = {"word/document.xml", "word/styles.xml", rel_name, ct}
        for name in original.namelist():
            if name not in editable:
                same = parts[name] == original.read(name)
                preserve_checks.append({"part": name, "unchanged": same})
                if not same:
                    raise AssertionError(f"Preserve-only part changed: {name}")
    FINAL.parent.mkdir(parents=True, exist_ok=True)
    with ZipFile(FINAL, "w", compression=ZIP_DEFLATED) as output:
        for name, data in parts.items():
            output.writestr(name, data)
    QA.mkdir(parents=True, exist_ok=True)
    check = Document(FINAL)
    assert len(check.sections) == 1
    assert digest(REFERENCE.read_bytes()) == REFERENCE_HASH
    all_text = "\n".join(p.text for p in check.paragraphs)
    assert all(label in all_text for label in ["一、", "二、", "三、", "四、", "五、", "六、", "七、", "八、"])
    (QA / "report-build-audit.json").write_text(json.dumps({
        "template_sha256": REFERENCE_HASH, "output_sha256": digest(FINAL.read_bytes()),
        "reference_unchanged": True, "section_count": 1, "manual_page_markers": page_markers,
        "paragraph_count": len(check.paragraphs), "image_occurrences": len(figures), "images": figures, "tables": table_audit,
        "typography": {"body": "宋体 + Times New Roman 12pt", "body_line_pt": 20, "first_line_indent_pt": 24,
                       "heading": "黑体 14pt bold", "subheading": "黑体 12pt bold", "caption_pt": 11,
                       "reference": "old/homework/作业报告1.pdf"},
        "preserve_only_parts": preserve_checks,
    }, ensure_ascii=False, indent=2), encoding="utf-8")
    print(f"Created {FINAL}; {len(figures)} PNG occurrences; {len(table_audit)} tables.")


if __name__ == "__main__":
    build()
