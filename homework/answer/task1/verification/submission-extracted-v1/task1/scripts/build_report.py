"""Build the report from the retained teacher template and task-local Markdown.

Run with the bundled Codex Python. The source template stays byte-for-byte intact.
Only the body, added styles, image relationships, and PNG content types may change.
"""
from __future__ import annotations

from copy import deepcopy
from hashlib import sha256
from io import BytesIO
import json
from pathlib import Path
import re
from zipfile import ZipFile, ZIP_DEFLATED

from docx import Document
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor
from docx.text.paragraph import Paragraph
from lxml import etree
from PIL import Image


TASK = Path(__file__).resolve().parents[1]
REPO = TASK.parents[2]
REFERENCE = REPO / "homework/teacher/task1/作业报告.docx"
FINAL = TASK / "docs/作业报告.docx"
MARKDOWN = TASK / "docs/报告.md"
QA = TASK / "verification/report"
REFERENCE_HASH = "cd682fbbda36cce208e9661699b5353fe4429ba343daa9787da6a6793ad24e16"
FONT = "Kaiti TC"


def digest(data: bytes) -> str:
    return sha256(data).hexdigest()


def append_xml(original: bytes, children: list[bytes], closing: bytes) -> bytes:
    if not children:
        return original
    if original.count(closing) != 1:
        raise ValueError(f"Ambiguous XML closing marker {closing!r}")
    return original.replace(closing, b"".join(children) + closing)


def build() -> None:
    if digest(REFERENCE.read_bytes()) != REFERENCE_HASH:
        raise ValueError("The retained report template hash changed.")
    text = MARKDOWN.read_text(encoding="utf-8")
    unresolved = re.findall(r"<!-- (?:robot|pixel-scene|texture|validation)-description -->", text)
    if unresolved:
        raise ValueError(f"Report evidence slots remain unresolved: {unresolved}")
    reference_doc = Document(REFERENCE)
    source_paragraphs = reference_doc.paragraphs
    prototypes = {
        "title": deepcopy(source_paragraphs[1]._p),
        "subtitle": deepcopy(source_paragraphs[2]._p),
        "author": deepcopy(source_paragraphs[3]._p),
        "heading": deepcopy(source_paragraphs[5]._p),
        "body": deepcopy(source_paragraphs[9]._p),
    }
    body_run_properties = deepcopy(source_paragraphs[9].runs[0]._r.rPr)
    source_section = deepcopy(reference_doc.sections[0]._sectPr)
    body = reference_doc._element.body
    for element in list(body):
        if element.tag != qn("w:sectPr"):
            body.remove(element)

    for name, size, title in [("Title", 24, True), ("ReportHeading", 14, False),
                              ("ReportCaption", 12, False)]:
        style = reference_doc.styles.add_style(name, WD_STYLE_TYPE.PARAGRAPH)
        style.base_style = reference_doc.styles["Normal"]
        style.font.name = FONT
        style.font.size = Pt(size)
        style.font.color.rgb = RGBColor(0, 0, 0)
        style.font.bold = title
        style.element.get_or_add_rPr().get_or_add_rFonts().set(qn("w:eastAsia"), FONT)
        style.paragraph_format.keep_with_next = name != "ReportCaption"

    pending_break = False
    figures = []
    page_markers = 0
    body_width = (reference_doc.sections[0].page_width - reference_doc.sections[0].left_margin
                  - reference_doc.sections[0].right_margin) / 914400.0

    def paragraph(content: str = "", role: str = "body", caption: bool = False) -> Paragraph:
        nonlocal pending_break
        element = deepcopy(prototypes[role])
        for child in list(element):
            if child.tag != qn("w:pPr"):
                element.remove(child)
        body.insert(len(body) - 1, element)
        p = Paragraph(element, reference_doc._body)
        if pending_break:
            p.paragraph_format.page_break_before = True
            pending_break = False
        if role == "title":
            p.style = reference_doc.styles["Title"]
        elif role == "heading":
            p.style = reference_doc.styles["ReportHeading"]
            p.paragraph_format.keep_with_next = True
            p.paragraph_format.space_after = Pt(8)
        elif caption:
            p.style = reference_doc.styles["ReportCaption"]
            p.paragraph_format.space_after = Pt(6)
        else:
            p.paragraph_format.space_after = Pt(6)
        if content:
            run = p.add_run(content)
            if role in ("title", "subtitle", "author"):
                source = source_paragraphs[{"title": 1, "subtitle": 2, "author": 3}[role]].runs[0]
                rpr = deepcopy(source._r.rPr)
            else:
                rpr = deepcopy(body_run_properties)
            if run._r.rPr is not None:
                run._r.remove(run._r.rPr)
            run._r.insert(0, rpr)
            for node in run._r.rPr.findall(qn("w:color")):
                run._r.rPr.remove(node)
            run.font.color.rgb = RGBColor(0, 0, 0)
            if caption:
                run.font.size = Pt(12)
            if role == "heading":
                run.font.bold = False
        return p

    lines = text.splitlines()
    buffer = []
    title_passed = False
    title_blocks = 0

    def flush() -> None:
        nonlocal title_blocks
        if not buffer:
            return
        content = "".join(buffer)
        buffer.clear()
        if title_passed and title_blocks < 2:
            paragraph(content, ["subtitle", "author"][title_blocks])
            title_blocks += 1
            return
        paragraph(content, caption=bool(re.match(r"图 \d+\.\d+", content)))

    for line in lines:
        if not line.strip():
            flush()
            continue
        if line.startswith("# "):
            flush()
            paragraph(line[2:], "title")
            title_passed = True
            continue
        if line.startswith("## "):
            flush()
            paragraph(line[3:], "heading")
            continue
        if line == "<!-- page -->":
            flush()
            pending_break = True
            page_markers += 1
            continue
        match = re.fullmatch(r"<!-- images: (.+) -->", line)
        if match:
            flush()
            names = [item.strip() for item in match.group(1).split("|")]
            p = paragraph()
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p.paragraph_format.keep_with_next = True
            p.paragraph_format.space_after = Pt(3)
            width = (body_width - 0.10 * (len(names) - 1)) / len(names)
            for i, name in enumerate(names):
                path = TASK / "docs/screenshots" / name
                if not path.is_file():
                    raise FileNotFoundError(path)
                with Image.open(path) as img:
                    if img.format != "PNG":
                        raise ValueError(f"Expected a lossless PNG: {path}")
                    pixels = img.size
                run = p.add_run()
                shape = run.add_picture(str(path), width=Inches(width))
                shape._inline.docPr.set("descr", name)
                figures.append({"file": str(path.relative_to(TASK)), "pixels": pixels,
                                "width_inches": width, "sha256": digest(path.read_bytes())})
                if i + 1 < len(names):
                    p.add_run(" ")
            continue
        if line.startswith("<!--"):
            flush()
            continue
        buffer.append(line)
    flush()

    # Confirm no geometry changed in the object model.
    def xml_values(element):
        return (element.tag, tuple(sorted(element.attrib.items())),
                tuple(xml_values(child) for child in element))
    if xml_values(source_section) != xml_values(reference_doc.sections[0]._sectPr):
        raise AssertionError("Report section geometry changed.")
    generated_stream = BytesIO()
    reference_doc.save(generated_stream)
    generated_stream.seek(0)
    preserve_checks = []
    with ZipFile(REFERENCE) as original, ZipFile(generated_stream) as generated:
        parts = {name: original.read(name) for name in original.namelist()}
        parts["word/document.xml"] = generated.read("word/document.xml")
        # Append only the three documented new styles. All old style bytes stay intact.
        generated_styles = etree.fromstring(generated.read("word/styles.xml"))
        new_style_nodes = [etree.tostring(node) for node in generated_styles
                           if node.get(qn("w:styleId")) in {"Title", "ReportHeading", "ReportCaption"}]
        parts["word/styles.xml"] = append_xml(parts["word/styles.xml"], new_style_nodes, b"</w:styles>")
        rel_name = "word/_rels/document.xml.rels"
        original_rels = etree.fromstring(original.read(rel_name))
        old_rel_ids = {node.get("Id") for node in original_rels}
        added_rels = [etree.tostring(node) for node in etree.fromstring(generated.read(rel_name))
                      if node.get("Id") not in old_rel_ids]
        parts[rel_name] = append_xml(parts[rel_name], added_rels, b"</Relationships>")
        content_name = "[Content_Types].xml"
        original_types = etree.fromstring(original.read(content_name))
        old_type_keys = {(node.tag, node.get("Extension"), node.get("PartName")) for node in original_types}
        added_types = [etree.tostring(node) for node in etree.fromstring(generated.read(content_name))
                       if (node.tag, node.get("Extension"), node.get("PartName")) not in old_type_keys]
        parts[content_name] = append_xml(parts[content_name], added_types, b"</Types>")
        for name in generated.namelist():
            if name.startswith("word/media/"):
                parts[name] = generated.read(name)
        editable = {"word/document.xml", "word/styles.xml", rel_name, content_name}
        for name in original.namelist():
            if name not in editable:
                preserve_checks.append({"part": name, "unchanged": parts[name] == original.read(name)})
                if not preserve_checks[-1]["unchanged"]:
                    raise AssertionError(f"Preserve-only part changed: {name}")
    FINAL.parent.mkdir(parents=True, exist_ok=True)
    with ZipFile(FINAL, "w", compression=ZIP_DEFLATED) as final_zip:
        for name, data in parts.items():
            final_zip.writestr(name, data)
    QA.mkdir(parents=True, exist_ok=True)
    check_doc = Document(FINAL)
    assert len(check_doc.sections) == 1
    assert digest(REFERENCE.read_bytes()) == REFERENCE_HASH
    text_all = "\n".join(p.text for p in check_doc.paragraphs)
    assert all(label in text_all for label in ["一、", "二、", "三、", "四、", "五、", "六、", "七、", "八、"])
    (QA / "report-build-audit.json").write_text(json.dumps({
        "template_sha256": REFERENCE_HASH,
        "output_sha256": digest(FINAL.read_bytes()),
        "reference_unchanged": True,
        "section_count": len(check_doc.sections),
        "manual_page_markers": page_markers,
        "paragraph_count": len(check_doc.paragraphs),
        "image_occurrences": len(figures),
        "images": figures,
        "preserve_only_parts": preserve_checks,
    }, ensure_ascii=False, indent=2), encoding="utf-8")
    print(f"Created {FINAL}; {len(figures)} PNG occurrences; all preserve-only parts unchanged.")


if __name__ == "__main__":
    build()
